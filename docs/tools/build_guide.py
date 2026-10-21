"""Build the standalone Word and HTML guides from the numbered Markdown chapters.

Run from any directory: python docs/tools/build_guide.py
Requires python-docx. No network, site server, or JavaScript dependencies are used.
"""

from __future__ import annotations

import html
import re
from pathlib import Path
from urllib.parse import unquote

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DOCS = Path(__file__).resolve().parents[1]
WORD_PATH = DOCS / 'Руководство по магазину.docx'
HTML_PATH = DOCS / 'Читать-в-браузере.html'
INLINE = re.compile(r'(`[^`\n]+`|\*\*.+?\*\*|\[[^\]]+\]\([^)]+\))')
CHAPTERS = sorted(DOCS.glob('[0-2][0-9]-*.md'))
CHAPTERS += [DOCS / 'Требования-к-фотографиям.md']
ANCHORS = {p.name: f'chapter_{i:02d}' for i, p in enumerate(CHAPTERS)}


def blocks(source):
    lines = source.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip() or re.fullmatch(r'\s*([-*_])\1{2,}\s*', line):
            i += 1
            continue
        if line.lstrip().startswith('```'):
            language = line.strip()[3:]
            code = []
            i += 1
            while i < len(lines) and not lines[i].lstrip().startswith('```'):
                code.append(lines[i])
                i += 1
            yield {'kind': 'code', 'text': '\n'.join(code), 'language': language}
            i += 1
            continue
        heading = re.match(r'^(#{1,6})\s+(.+)', line)
        if heading:
            yield {'kind': 'heading', 'level': len(heading[1]), 'text': heading[2]}
            i += 1
            continue
        if line.lstrip().startswith('|'):
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith('|'):
                cells = [c.strip().replace('\\|', '|') for c in re.split(r'(?<!\\)\|', lines[i].strip().strip('|'))]
                if not all(re.fullmatch(r':?-{2,}:?', c.replace(' ', '')) for c in cells):
                    rows.append(cells)
                i += 1
            yield {'kind': 'table', 'rows': rows}
            continue
        item = re.match(r'^(\s*)([-*]|\d+[.)])\s+(.+)', line)
        if item:
            text = item[3]
            depth = len(item[1]) // 2
            i += 1
            while i < len(lines) and lines[i].strip() and re.match(r'^\s{2,}\S', lines[i]) and not re.match(r'^\s*([-*]|\d+[.)])\s+', lines[i]):
                text += ' ' + lines[i].strip()
                i += 1
            yield {'kind': 'list', 'text': text, 'prefix': item[2], 'depth': depth}
            continue
        parts = [line.lstrip('> ').strip() if line.startswith('>') else line.strip()]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r'^\s*(#{1,6}\s|```|\||[-*]\s|\d+[.)]\s|>)', lines[i]):
            parts.append(lines[i].strip())
            i += 1
        yield {'kind': 'paragraph', 'text': ' '.join(parts)}


def plain(text):
    text = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', text)
    return text.replace('**', '').replace('`', '')


def heading_text(text):
    # Headings in the reading editions use plain words; exact identifiers stay in body text.
    return re.sub(r'\s+', ' ', re.sub(r'[^\w\s]', ' ', plain(text))).strip()


def bookmark(paragraph, name, number):
    start = OxmlElement('w:bookmarkStart')
    start.set(qn('w:id'), str(number))
    start.set(qn('w:name'), name)
    end = OxmlElement('w:bookmarkEnd')
    end.set(qn('w:id'), str(number))
    paragraph._p.insert(1 if paragraph._p.pPr is not None else 0, start)
    paragraph._p.append(end)


def link_target(url):
    name = unquote(url.split('#')[0]).replace('\\', '/').split('/')[-1]
    if name in ANCHORS:
        return '#' + ANCHORS[name]
    if name == HTML_PATH.name:
        return HTML_PATH.name
    if name == WORD_PATH.name:
        return WORD_PATH.name
    if url.startswith(('https://', 'http://', 'mailto:')):
        return url
    return None


def add_inline(paragraph, text):
    for token in INLINE.split(text):
        if not token:
            continue
        match = re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)', token)
        if match:
            label, url = match.groups()
            target = link_target(url)
            if target and target.startswith('#'):
                hyperlink = OxmlElement('w:hyperlink')
                hyperlink.set(qn('w:anchor'), target[1:])
                hyperlink.set(qn('w:history'), '1')
                run = OxmlElement('w:r')
                props = OxmlElement('w:rPr')
                color = OxmlElement('w:color')
                color.set(qn('w:val'), '24556B')
                props.append(color)
                underline = OxmlElement('w:u')
                underline.set(qn('w:val'), 'single')
                props.append(underline)
                run.append(props)
                text_node = OxmlElement('w:t')
                text_node.text = plain(label)
                run.append(text_node)
                hyperlink.append(run)
                paragraph._p.append(hyperlink)
            elif target and target.startswith(('https://', 'http://', 'mailto:')):
                hyperlink = OxmlElement('w:hyperlink')
                rel = paragraph.part.relate_to(target, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
                hyperlink.set(qn('r:id'), rel)
                run = OxmlElement('w:r')
                props = OxmlElement('w:rPr')
                color = OxmlElement('w:color')
                color.set(qn('w:val'), '24556B')
                props.append(color)
                run.append(props)
                node = OxmlElement('w:t')
                node.text = plain(label)
                run.append(node)
                hyperlink.append(run)
                paragraph._p.append(hyperlink)
            else:
                paragraph.add_run(plain(label))
        elif token.startswith('`') and token.endswith('`'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
        elif token.startswith('**') and token.endswith('**'):
            run_count = len(paragraph._p.findall('.//' + qn('w:r')))
            add_inline(paragraph, token[2:-2])
            for node in paragraph._p.findall('.//' + qn('w:r'))[run_count:]:
                props = node.get_or_add_rPr()
                if props.find(qn('w:b')) is None:
                    props.append(OxmlElement('w:b'))
        else:
            paragraph.add_run(token)


def style_document(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin, sec.bottom_margin = Cm(1.65), Cm(1.65)
    sec.footer_distance = Cm(.7)
    for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Heading 4', 'List Paragraph'):
        s = doc.styles[name]
        s.font.name = 'Arial'
        s.font.color.rgb = RGBColor(0, 0, 0)
        fonts = s.element.get_or_add_rPr().get_or_add_rFonts()
        fonts.set(qn('w:eastAsia'), 'Arial')
        for attr in list(fonts.attrib):
            if attr.lower().endswith('theme'):
                del fonts.attrib[attr]
        for border in list(s.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
        lang = OxmlElement('w:lang')
        lang.set(qn('w:val'), 'ru-RU')
        s.element.rPr.append(lang)
    normal = doc.styles['Normal']
    normal.font.size = Pt(11.5)
    normal.paragraph_format.line_spacing = 1.1
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.widow_control = True
    doc.styles['Title'].font.size = Pt(26)
    doc.styles['Title'].font.bold = True
    doc.styles['Title'].paragraph_format.space_after = Pt(15)
    doc.styles['Subtitle'].font.size = Pt(13)
    doc.styles['Subtitle'].font.italic = False
    for spacing in list(doc.styles['Subtitle'].element.rPr.findall(qn('w:spacing'))):
        doc.styles['Subtitle'].element.rPr.remove(spacing)
    for level, size in ((1, 21), (2, 15), (3, 12.5), (4, 11.5)):
        s = doc.styles[f'Heading {level}']
        s.font.size = Pt(size)
        s.font.bold = True
        s.paragraph_format.space_before = Pt(14 if level > 1 else 0)
        s.paragraph_format.space_after = Pt(7)
        s.paragraph_format.keep_with_next = True
        s.paragraph_format.keep_together = True
    code = doc.styles.add_style('Guide Code', 1)
    code.font.name = 'Consolas'
    code.font.size = Pt(9.2)
    code.font.color.rgb = RGBColor(0, 0, 0)
    code.paragraph_format.line_spacing = 1.06
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.left_indent = Cm(.22)
    footer = sec.footer.paragraphs[0]
    footer.alignment = 2
    run = footer.add_run()
    run.font.name, run.font.size = 'Arial', Pt(9)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    run._r.addnext(field)
    doc.core_properties.title = 'Руководство по магазину K Beauty C Health'
    doc.core_properties.subject = 'Запуск управление и устройство магазина простыми словами'
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    doc.core_properties.comments = ''


def word_table(doc, rows):
    count = max(map(len, rows))
    rows = [row + [''] * (count - len(row)) for row in rows]
    # Long reference records read better on a phone as label and value paragraphs.
    if count > 3 or any(max((len(plain(c)) for c in row), default=0) > 215 for row in rows[1:]):
        headers = rows[0]
        for row in rows[1:]:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.keep_with_next = True
            add_inline(p, '**' + plain(row[0]) + '**')
            for label, value in zip(headers[1:], row[1:]):
                p = doc.add_paragraph()
                add_inline(p, '**' + plain(label) + ':** ' + value)
        return
    table = doc.add_table(rows=1, cols=count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    weights = [max(12, min(72, sum(len(plain(row[c])) for row in rows) / len(rows))) for c in range(count)]
    total = sum(weights)
    widths = [17.4 * value / total for value in weights]
    widths = [max(3.0 if count == 3 else 4.5, w) for w in widths]
    widths = [w * 17.4 / sum(widths) for w in widths]
    for c, width in zip(table.columns, widths):
        c.width = Cm(width)
    for index, row in enumerate(rows):
        cells = table.rows[0].cells if index == 0 else table.add_row().cells
        table.rows[index]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for col, (cell, value) in enumerate(zip(cells, row)):
            cell.width = Cm(widths[col])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side in ('top', 'left', 'bottom', 'right'):
                node = OxmlElement('w:' + side)
                node.set(qn('w:w'), '90')
                node.set(qn('w:type'), 'dxa')
                margins.append(node)
            tcpr.append(margins)
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'E8EDF0' if index == 0 else ('F6F7F8' if index % 2 == 0 else 'FFFFFF'))
            tcpr.append(shade)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.line_spacing = 1.08
            p.paragraph_format.keep_with_next = index == 0
            add_inline(p, value)
            for run in p.runs:
                run.font.size = Pt(10)
                if index == 0:
                    run.bold = True
        if index == 0:
            repeat = OxmlElement('w:tblHeader')
            table.rows[0]._tr.get_or_add_trPr().append(repeat)
    borders = OxmlElement('w:tblBorders')
    for name in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge = OxmlElement('w:' + name)
        edge.set(qn('w:val'), 'single')
        edge.set(qn('w:sz'), '4')
        edge.set(qn('w:color'), 'D9D9D9')
        borders.append(edge)
    table._tbl.tblPr.append(borders)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(3)
    spacer.add_run().font.size = Pt(3)


def make_word(parsed):
    doc = Document()
    style_document(doc)
    for chapter_index, (path, chapter_blocks) in enumerate(parsed):
        for b in chapter_blocks:
            kind = b['kind']
            if kind == 'heading':
                if b['level'] == 1:
                    if chapter_index == 0:
                        p = doc.add_paragraph('Руководство по магазину\nK Beauty C Health', 'Title')
                        sub = doc.add_paragraph('Понятная инструкция для владельца и начинающего разработчика', 'Subtitle')
                        sub.paragraph_format.space_after = Pt(18)
                    else:
                        p = doc.add_paragraph(heading_text(b['text']), 'Heading 1')
                        p.paragraph_format.page_break_before = chapter_index == 1
                        p.paragraph_format.space_before = Pt(22)
                        p.paragraph_format.space_after = Pt(10)
                    bookmark(p, ANCHORS[path.name], chapter_index + 1)
                else:
                    doc.add_paragraph(heading_text(b['text']), f'Heading {min(b["level"], 4)}')
            elif kind == 'paragraph':
                p = doc.add_paragraph()
                add_inline(p, b['text'])
            elif kind == 'list':
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(.55 + b['depth'] * .35)
                p.paragraph_format.first_line_indent = Cm(-.45)
                p.paragraph_format.space_after = Pt(4)
                prefix = '•' if b['prefix'] in ('-', '*') else b['prefix']
                p.add_run(prefix + '  ')
                add_inline(p, b['text'])
            elif kind == 'table' and b['rows']:
                word_table(doc, b['rows'])
            elif kind == 'code':
                code_lines = b['text'].split('\n')
                for line_index, line in enumerate(code_lines):
                    p = doc.add_paragraph(line or ' ', 'Guide Code')
                    p.paragraph_format.space_before = Pt(4 if line_index == 0 else 0)
                    p.paragraph_format.space_after = Pt(7 if line_index == len(code_lines) - 1 else 0)
                    p.paragraph_format.keep_with_next = line_index < len(code_lines) - 1 and len(code_lines) <= 12
                    p.paragraph_format.keep_together = True
    doc.save(WORD_PATH)


def inline_html(text):
    parts = []
    for token in INLINE.split(text):
        match = re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)', token)
        if match:
            label, url = match.groups()
            target = link_target(url)
            parts.append(f'<a href="{html.escape(target, quote=True)}">{html.escape(plain(label))}</a>' if target else html.escape(plain(label)))
        elif token.startswith('`') and token.endswith('`'):
            parts.append('<code>' + html.escape(token[1:-1]) + '</code>')
        elif token.startswith('**') and token.endswith('**'):
            parts.append('<strong>' + inline_html(token[2:-2]) + '</strong>')
        else:
            parts.append(html.escape(token))
    return ''.join(parts)


CSS = '''
:root{color-scheme:light;--ink:#172127;--muted:#53616a;--line:#d9dfe3;--link:#24556b}
*{box-sizing:border-box}html{scroll-behavior:smooth;scroll-padding-top:4rem}
body{margin:0;background:#f4f5f6;color:var(--ink);font:18px/1.65 Arial,Helvetica,sans-serif}
main{max-width:900px;margin:auto;background:white;padding:36px 52px 80px}
.topbar{position:sticky;top:0;z-index:2;background:#fff;border-bottom:1px solid var(--line);padding:10px 18px;display:flex;gap:24px;justify-content:center;font-size:15px}
a{color:var(--link);text-decoration-thickness:1px;text-underline-offset:3px}a:focus-visible{outline:3px solid #9ac7d9;outline-offset:4px}
h1,h2,h3,h4{color:#000;line-height:1.22;text-wrap:balance}h1{font-size:2rem;margin:0 0 1rem}h2{font-size:1.55rem;margin:2.2rem 0 1rem}h3{font-size:1.22rem;margin:1.7rem 0 .7rem}h4{font-size:1.05rem}
section{padding-top:30px;margin-top:30px;border-top:1px solid var(--line)}section:first-child{margin-top:0;border:0;padding-top:0}
p{margin:.7rem 0 1rem}li{margin:.35rem 0}ul,ol{padding-left:1.6rem}
code{font: .88em/1.5 Consolas,'Courier New',monospace;overflow-wrap:anywhere}pre{background:#f5f7f8;padding:16px;white-space:pre-wrap;overflow-wrap:anywhere;border-radius:4px;font-size:15px;line-height:1.5}
table{border-collapse:collapse;width:100%;margin:1rem 0 1.7rem;font-size:16px;line-height:1.45;table-layout:fixed}
th,td{border:1px solid #d9d9d9;padding:11px 12px;text-align:left;vertical-align:middle;overflow-wrap:anywhere}th{background:#e8edf0;color:#000}tbody tr:nth-child(even){background:#f6f7f8}
.sub{font-size:1.1rem;color:var(--muted);margin:1rem 0 2rem}.back{font-size:14px;margin:1.8rem 0 .5rem}
@media(max-width:640px){body{font-size:17px;line-height:1.62;background:#fff}main{padding:24px 19px 56px}h1{font-size:1.75rem}h2{font-size:1.4rem}.topbar{font-size:14px;padding:10px;gap:22px}pre{padding:12px;font-size:13px}table,thead,tbody,tr,td{display:block;width:100%}thead{position:absolute;width:1px;height:1px;overflow:hidden;clip-path:inset(50%)}tr{margin:12px 0;border:1px solid var(--line)}td{border:0;border-top:1px solid #e5e9ec;padding:9px 12px}td:first-child{border:0;font-weight:bold}td::before{content:attr(data-label);display:block;font-weight:bold;font-size:.84em;color:var(--muted);margin-bottom:3px}td:first-child::before{display:none}tbody tr:nth-child(even){background:#fafbfc}}
@media print{body{background:white;font-size:11pt}main{max-width:none;padding:0}.topbar,.back{display:none}section{break-before:page;border:0}section:first-child{break-before:auto}h1,h2,h3{break-after:avoid}pre{font-size:9pt}a{color:inherit}table{font-size:9pt}tr{break-inside:avoid}}
'''


def make_html(parsed):
    out = ['<!doctype html><html lang="ru"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Руководство по магазину K Beauty C Health</title><style>' + CSS + '</style></head><body>',
           '<nav class="topbar" aria-label="Навигация"><a href="#contents">Оглавление</a><a href="Руководство%20по%20магазину.docx">Открыть Word</a></nav><main>']
    for chapter_index, (path, chapter_blocks) in enumerate(parsed):
        out.append(f'<section id="{ANCHORS[path.name]}">')
        list_tag = None
        for b in chapter_blocks:
            if b['kind'] != 'list' and list_tag:
                out.append(f'</{list_tag}>')
                list_tag = None
            if b['kind'] == 'heading':
                level = min(b['level'] + (1 if chapter_index else 0), 5)
                anchor = ' id="contents"' if chapter_index == 0 and plain(b['text']) == 'Все главы' else ''
                out.append(f'<h{level}{anchor}>{html.escape(heading_text(b["text"]))}</h{level}>')
                if chapter_index == 0 and b['level'] == 1:
                    out.append('<p class="sub">Понятная инструкция для владельца и начинающего разработчика</p>')
            elif b['kind'] == 'paragraph':
                out.append('<p>' + inline_html(b['text']) + '</p>')
            elif b['kind'] == 'list':
                tag = 'ul' if b['prefix'] in ('-', '*') else 'ol'
                if list_tag != tag:
                    if list_tag:
                        out.append(f'</{list_tag}>')
                    start = f' start="{re.sub(chr(92) + "D", "", b["prefix"])}"' if tag == 'ol' else ''
                    out.append(f'<{tag}{start}>')
                    list_tag = tag
                out.append('<li>' + inline_html(b['text']) + '</li>')
            elif b['kind'] == 'code':
                out.append('<pre><code>' + html.escape(b['text']) + '</code></pre>')
            elif b['kind'] == 'table' and b['rows']:
                rows = b['rows']
                out.append('<table><thead><tr>' + ''.join('<th scope="col">' + inline_html(c) + '</th>' for c in rows[0]) + '</tr></thead><tbody>')
                for row in rows[1:]:
                    out.append('<tr>' + ''.join(f'<td data-label="{html.escape(plain(rows[0][i]) if i < len(rows[0]) else "", quote=True)}">{inline_html(c)}</td>' for i, c in enumerate(row)) + '</tr>')
                out.append('</tbody></table>')
        if list_tag:
            out.append(f'</{list_tag}>')
        if chapter_index:
            out.append('<p class="back"><a href="#contents">Вернуться к оглавлению</a></p>')
        out.append('</section>')
    out.append('</main></body></html>')
    HTML_PATH.write_text('\n'.join(out), encoding='utf-8')


def main():
    parsed = [(path, list(blocks(path.read_text(encoding='utf-8-sig')))) for path in CHAPTERS]
    make_word(parsed)
    make_html(parsed)
    words = sum(len(re.findall(r'\S+', path.read_text(encoding='utf-8-sig'))) for path in CHAPTERS)
    print(f'Created {WORD_PATH.name} and {HTML_PATH.name}; {len(parsed)} source files; {words} source words')


if __name__ == '__main__':
    main()
